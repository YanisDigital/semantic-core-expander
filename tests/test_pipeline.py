"""Collector + exporter end to end with a fake fetcher (no network)."""
import contextlib
import csv
import io
import json
import sqlite3
import unittest

from helpers import FakeFetcher, make_project
import clean_export
import suggest_collect
from suggest_collect import Blocked, parse_response

FULL = [f"dentist {w}" for w in ["near me", "salary", "office", "cost", "open now",
                                  "school", "jobs", "reviews", "hours", "for kids",
                                  "emergency", "cleaning", "insurance", "visit", "chair"]]
ANSWERS = {
    "dentist": FULL,
    "dentist near me ": ["dentist near me open now", "dentist near me cheap",
                        "=HYPERLINK(\"http://x\") dentist"],
    "dentist cost": ["dentist cost without insurance", "cost dentist", "dental cost"],
    "dentist j": ["dentist jobs", "dentist jobs remote"],
}


def run_collect(project, fetcher, *args):
    with contextlib.redirect_stdout(io.StringIO()):
        return suggest_collect.main(["--project", str(project), *args], fetcher=fetcher)


def run_export(project, *args):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        code = clean_export.main(["--project", str(project), *args])
    return code, buf.getvalue()


class ParseResponseTest(unittest.TestCase):
    def test_shape(self):
        r = parse_response(["q", ["a", 5, "b" * 500], [], {}, {"google:suggestrelevance": [1, "x"]}])
        self.assertEqual(r["s"][0], "a")
        self.assertEqual(len(r["s"][1]), 200)
        self.assertEqual(r["rel"], [1])
        with self.assertRaises(ValueError):
            parse_response({"evil": True})


class PipelineTest(unittest.TestCase):
    def setUp(self):
        self.proj = make_project({"modifiers": ["cost"]})

    def test_collect_and_export(self):
        fetcher = FakeFetcher(ANSWERS)
        self.assertEqual(run_collect(self.proj, fetcher), 0)
        self.assertIn("dentist", fetcher.calls)
        # Full list -> one letter deeper; suggestion -> new prefix with a space.
        self.assertIn("dentist j", fetcher.calls)
        self.assertIn("dentist near me ", fetcher.calls)

        status = json.loads((self.proj / "data" / "status.json").read_text(encoding="utf-8"))
        self.assertEqual(status["state"], "done")
        con = sqlite3.connect(self.proj / "data" / "cache.sqlite")
        keys = [k for (k,) in con.execute("SELECT key FROM cache")]
        con.close()
        self.assertTrue(all(k.startswith("us|en|") for k in keys))

        # Re-run replays the cache: no new requests.
        again = FakeFetcher(ANSWERS)
        self.assertEqual(run_collect(self.proj, again), 0)
        self.assertEqual(again.calls, [])

        code, out = run_export(self.proj, "--preview", "5")
        self.assertEqual(code, 0)
        self.assertIn("Kept:", out)
        self.assertFalse((self.proj / "output").exists())

        code, _ = run_export(self.proj)
        self.assertEqual(code, 0)
        with open(self.proj / "output" / "test_keywords.csv", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
        kws = [r[0] for r in rows[1:]]
        self.assertIn("dentist jobs remote", kws)
        self.assertNotIn("dental cost", kws)                  # no root
        self.assertEqual(sum(k in ("dentist cost", "cost dentist") for k in kws), 1)  # word order
        evil = [k for k in kws if "hyperlink" in k]
        self.assertTrue(evil and all(k.startswith("'") for k in evil))
        for r in rows[1:]:
            self.assertTrue(1 <= int(r[1]) <= 100)

    def test_xlsx_has_text_not_formulas(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl not installed")
        run_collect(self.proj, FakeFetcher(ANSWERS))
        run_export(self.proj)
        wb = load_workbook(self.proj / "output" / "test_semantic_core.xlsx")
        self.assertEqual(wb.sheetnames, ["Keywords", "Clusters", "Rejected", "About"])
        cells = [c for row in wb["Keywords"].iter_rows(min_row=2) for c in row]
        evil = [c for c in cells if isinstance(c.value, str) and c.value.startswith("=")]
        self.assertTrue(evil)
        self.assertTrue(all(c.data_type == "s" for c in evil))

    def test_budget_and_resume(self):
        self.assertEqual(run_collect(self.proj, FakeFetcher(ANSWERS), "--max-requests", "3"), 0)
        status = json.loads((self.proj / "data" / "status.json").read_text(encoding="utf-8"))
        self.assertEqual(status["state"], "budget_reached")
        self.assertEqual(status["network"], 3)
        self.assertGreater(status["queue"], 0)

    def test_block_stops_with_code_3(self):
        fetcher = FakeFetcher(ANSWERS, fail_after=2, exc=Blocked(429))
        self.assertEqual(run_collect(self.proj, fetcher), 3)
        status = json.loads((self.proj / "data" / "status.json").read_text(encoding="utf-8"))
        self.assertEqual(status["state"], "blocked")
        con = sqlite3.connect(self.proj / "data" / "cache.sqlite")
        self.assertEqual(con.execute("SELECT COUNT(*) FROM cache").fetchone()[0], 2)
        con.close()

    def test_config_error_code_2(self):
        (self.proj / "project.json").write_text('{"project": "x"}', encoding="utf-8")
        with contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(run_collect(self.proj, FakeFetcher({})), 2)

    def test_dry_run_makes_no_requests(self):
        fetcher = FakeFetcher(ANSWERS)
        self.assertEqual(run_collect(self.proj, fetcher, "--dry-run"), 0)
        self.assertEqual(fetcher.calls, [])

    def test_cluster_labels(self):
        run_collect(self.proj, FakeFetcher(ANSWERS))
        (self.proj / "cluster_labels.json").write_text(
            json.dumps({"labels": {"job": "Careers"}}), encoding="utf-8")
        run_export(self.proj)
        with open(self.proj / "output" / "test_keywords.csv", encoding="utf-8-sig") as f:
            topics = {r[0]: r[2] for r in csv.reader(f)}
        self.assertEqual(topics["dentist jobs remote"], "Careers")


if __name__ == "__main__":
    unittest.main()
