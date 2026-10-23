"""Clean the collected suggestions, score them, cluster them and export.

Reads <project>/data/cache.sqlite (filled by suggest_collect.py), computes a
relative popularity index, drops duplicates / foreign-language / off-topic
phrases, groups the rest into topic clusters and writes
<project>/output/<slug>_semantic_core.xlsx (+ CSV files).

Usage:
  python clean_export.py --project DIR            # write the files
  python clean_export.py --project DIR --preview  # print samples, write nothing

Optional <project>/cluster_labels.json renames / merges clusters:
  {"labels": {"stem": "Topic name", ...}, "ignore_stems": ["stem", ...]}
"""
import argparse
import csv
import json
import math
import random
import sqlite3
import sys
from collections import Counter, defaultdict
from datetime import date

from common import ConfigError, csv_safe, inside, load_project
from textutils import Filters, normalize, strip_controls, tokens

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

CORE_TOPIC = "(core)"
MISC_TOPIC = "(misc)"


def collect(cfg):
    """{phrase: raw score}, {phrase: hits} and the number of cached answers."""
    path = cfg["data_dir"] / "cache.sqlite"
    if not path.is_file():
        raise ConfigError(f"no cache at {path}; run suggest_collect.py first")
    con = sqlite3.connect(path)
    profile = cfg["profile"]
    roots = [normalize(r, profile) for r in cfg["roots"]]
    offset = max(0, min(len(r) for r in roots) - 1)
    score, hits, answers = defaultdict(float), defaultdict(int), 0
    rows = con.execute("SELECT q, resp FROM cache WHERE gl = ? AND hl = ?",
                       (cfg["gl"], cfg["hl"]))
    for q, resp in rows:
        answers += 1
        nq = normalize(q, profile)
        typed = max(1, len(q.strip()) - offset)  # characters typed beyond the root
        for pos, s in enumerate(json.loads(resp)["s"][:15]):
            s = normalize(s, profile)
            if not s:
                continue
            pos_w = (16 - pos) / 15
            # A prefix match is stronger evidence than an "infix" one.
            match_w = 1.0 if s.startswith(nq) else 0.6
            score[s] += pos_w * match_w / math.sqrt(typed)
            hits[s] += 1
    con.close()
    return score, hits, answers


def to_index(score):
    if not score:
        return {}
    vals = [math.log(v) for v in score.values()]
    lo, hi = min(vals), max(vals)
    span = (hi - lo) or 1
    return {k: max(1, round(1 + 99 * (math.log(v) - lo) / span)) for k, v in score.items()}


def clean(cfg, score):
    flt = Filters(cfg)
    kept, rejected = {}, []
    for s in score:
        if not flt.relevant(s):
            continue  # infix noise without the root is not part of the core
        r = flt.reject_reason(s)
        if r:
            rejected.append((s, r))
        else:
            kept[s] = score[s]
    if cfg["cleaning"]["merge_word_order"]:
        by_key = defaultdict(list)
        for s in kept:
            by_key[" ".join(sorted(tokens(s)))].append(s)
        for group in by_key.values():
            if len(group) > 1:
                group.sort(key=lambda x: (-kept[x], x))
                for d in group[1:]:
                    rejected.append((d, f"duplicate (word order of: {group[0]})"))
                    del kept[d]
    return kept, rejected


def load_labels(cfg):
    path = cfg["dir"] / "cluster_labels.json"
    if not path.is_file():
        return {}, set()
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    labels = data.get("labels") if isinstance(data, dict) else None
    ignore = data.get("ignore_stems") if isinstance(data, dict) else None
    if not isinstance(data, dict) or not isinstance(labels or {}, dict) \
            or not isinstance(ignore or [], list):
        raise ConfigError('cluster_labels.json must be {"labels": {stem: topic}, '
                          '"ignore_stems": [stem, ...]}')
    labels = {str(k): strip_controls(str(v)).strip()[:80] for k, v in (labels or {}).items()}
    return labels, {str(x) for x in ignore or []}


def cluster(cfg, kept, idx, labels=None, ignore=None):
    """Assign every keyword to a topic: its most frequent content stem."""
    p = cfg["profile"]
    labels, ignore = labels or {}, ignore or set()
    stem_len = int(p["stem_len"])
    suffixes = sorted(p.get("strip_suffixes") or [], key=len, reverse=True)
    roots = [normalize(r, p) for r in cfg["roots"]]
    skip = set(p["stopwords"]) | {normalize(w, p) for w in cfg["questions"]}
    for s in cfg["seeds"]:
        skip.update(tokens(normalize(s, p)))
    for r in roots:
        skip.update(tokens(r))

    def stems(s):
        out = []
        for t in tokens(s):
            if t in skip or t.isdigit() or len(t) < 2 or any(r in t for r in roots if " " not in r):
                continue
            for suf in suffixes:
                if t.endswith(suf) and len(t) - len(suf) >= 3 and not t.endswith(suf * 2):
                    t = t[:-len(suf)]
                    break
            st = t[:stem_len]
            if st not in ignore and st not in out:
                out.append(st)
        return out

    per_kw = {s: stems(s) for s in kept}
    df = Counter(st for sts in per_kw.values() for st in sts)
    topic_of = {}
    for s, sts in per_kw.items():
        topic_of[s] = max(sts, key=lambda st: (df[st], -sts.index(st))) if sts else CORE_TOPIC

    size = Counter(topic_of.values())
    groups = defaultdict(list)
    min_size = cfg["cleaning"]["min_cluster_size"]
    for s, st in topic_of.items():
        if st not in (CORE_TOPIC,) and size[st] < min_size and st not in labels:
            st = MISC_TOPIC
        name = labels.get(st, st)
        topic_of[s] = name
        groups[name].append((s, st))

    clusters = []
    for name, items in groups.items():
        items.sort(key=lambda x: (-idx[x[0]], x[0]))
        clusters.append({
            "topic": name,
            "stems": sorted({st for _, st in items if st not in (CORE_TOPIC, MISC_TOPIC)}),
            "queries": len(items),
            "sum_index": sum(idx[s] for s, _ in items),
            "max_index": max(idx[s] for s, _ in items),
            "top": [s for s, _ in items[:5]],
        })
    clusters.sort(key=lambda c: (-c["sum_index"], c["topic"]))
    return topic_of, clusters


def preview(cfg, kept, rejected, idx, clusters, answers, n):
    rnd = random.Random(42)
    rows = sorted(kept, key=lambda s: (-idx[s], s))
    print(f"Cached answers: {answers}. Kept: {len(kept)}, rejected: {len(rejected)}")
    for stem, hits in Filters(cfg).exclude_conflicts(cfg):
        print(f"Warning: exclude stem '{stem}' also matches your own words: "
              f"{', '.join(hits[:5])}")
    print("\n# Top by index")
    for s in rows[:n]:
        print(f"{idx[s]:>3}  {s}")
    print("\n# Random kept (check relevance / other meanings)")
    for s in rnd.sample(rows, min(n, len(rows))):
        print(f"{idx[s]:>3}  {s}")
    by_reason = defaultdict(list)
    for s, r in rejected:
        by_reason[r.split(" (")[0].split(":")[0]].append((s, r))
    print("\n# Rejected by reason")
    for reason, items in sorted(by_reason.items(), key=lambda x: -len(x[1])):
        print(f"- {reason}: {len(items)}")
        for s, r in rnd.sample(items, min(5, len(items))):
            print(f"    {s}   [{r}]")
    print("\n# Clusters (topic | queries | sum index | examples)")
    for c in clusters[:max(n, 40)]:
        print(f"{c['topic']} | {c['queries']} | {c['sum_index']} | {'; '.join(c['top'][:3])}")


def write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(header)
        for r in rows:
            w.writerow([csv_safe(v) for v in r])


def write_xlsx(path, sheets):
    try:
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return False
    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F5597")
    for title, header, rows, widths in sheets:
        ws = wb.create_sheet(title)
        for c, h in enumerate(header, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font, cell.fill = head_font, head_fill
        for r, row in enumerate(rows, 2):
            for c, v in enumerate(row, 1):
                if isinstance(v, str):
                    v = ILLEGAL_CHARACTERS_RE.sub("", v)
                cell = ws.cell(row=r, column=c, value=v)
                if isinstance(v, str):
                    cell.data_type = "s"  # never let a phrase become a formula
                    if title == "About":
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, wdt in enumerate(widths):
            ws.column_dimensions[chr(ord("A") + i)].width = wdt
    wb.save(path)
    return True


def main(argv=None):
    ap = argparse.ArgumentParser(description="Clean, cluster and export suggestions.")
    ap.add_argument("--project", required=True, help="project directory with project.json")
    ap.add_argument("--preview", nargs="?", type=int, const=30, metavar="N",
                    help="print N sample rows per section and write nothing")
    args = ap.parse_args(argv)

    try:
        cfg = load_project(args.project)
        score, hits, answers = collect(cfg)
        labels, ignore = load_labels(cfg)
    except (ConfigError, json.JSONDecodeError) as e:
        print(f"Error: {e}", file=sys.stderr)
        return 2

    kept, rejected = clean(cfg, score)
    idx = to_index(kept)
    topic_of, clusters = cluster(cfg, kept, idx, labels, ignore)
    if args.preview is not None:
        preview(cfg, kept, rejected, idx, clusters, answers, max(1, args.preview))
        return 0

    out_dir = inside(cfg["dir"], cfg["output_dir"])
    out_dir.mkdir(parents=True, exist_ok=True)
    slug = cfg["project"]
    rows = sorted(kept, key=lambda s: (-kept[s], s))
    kw_rows = [(s, idx[s], topic_of[s]) for s in rows]
    cl_rows = [(c["topic"], ", ".join(c["stems"]), c["queries"], c["sum_index"],
                c["max_index"], "; ".join(c["top"])) for c in clusters]
    rej_rows = sorted(rejected, key=lambda x: (x[1], x[0]))
    p = cfg["profile"]
    about = [
        ("Source", "Google Autocomplete (suggestqueries.google.com), unofficial endpoint"),
        ("Country / language", f"gl={cfg['gl']}, hl={cfg['hl']}, profile={cfg['language']} "
                               f"({p.get('name', cfg['language'])})"),
        ("Exported", date.today().isoformat()),
        ("Site / goal", " | ".join(x for x in (cfg["site"], cfg["goal"]) if x) or "-"),
        ("Seeds", "; ".join(cfg["seeds"])),
        ("Relevance roots", "; ".join(cfg["roots"])),
        ("Method", "Recursive crawl of autocomplete: alphabet expansion after every full "
                   "prefix, niche and question modifiers, every found phrase expanded again."),
        ("Index (1-100)", "RELATIVE popularity, not monthly searches. Higher when Google "
                          "suggests the phrase after a shorter prefix, at a higher position "
                          "and for more prefixes. Log scale, comparable only inside this file."),
        ("Cleaning", "Normalised case / apostrophes / spaces; dropped duplicates (incl. "
                     "word order), wrong script, excluded languages "
                     f"({', '.join(cfg['excluded']) or 'none'}), adult/profanity, project "
                     "excludes, truncated fragments, overlong and junk phrases. See 'Rejected'."),
        ("Clusters", "Grouped by the most frequent content word stem; small groups go to "
                     f"{MISC_TOPIC}, phrases with only the root to {CORE_TOPIC}."),
        ("Cached answers", answers),
        ("Keywords", len(rows)),
        ("Rejected", len(rejected)),
        ("Clusters total", len(clusters)),
    ]

    xlsx = out_dir / f"{slug}_semantic_core.xlsx"
    ok = write_xlsx(xlsx, [
        ("Keywords", ["Keyword", "Index (1-100)", "Cluster"], kw_rows, [70, 16, 28]),
        ("Clusters", ["Topic", "Stems", "Queries", "Sum index", "Max index", "Top keywords"],
         cl_rows, [28, 24, 10, 12, 12, 90]),
        ("Rejected", ["Keyword", "Reason"], rej_rows, [70, 60]),
        ("About", ["Parameter", "Value"], about, [22, 110]),
    ])
    write_csv(out_dir / f"{slug}_keywords.csv", ["keyword", "index", "cluster"], kw_rows)
    write_csv(out_dir / f"{slug}_clusters.csv",
              ["topic", "stems", "queries", "sum_index", "max_index", "top"], cl_rows)
    write_csv(inside(cfg["dir"], cfg["data_dir"] / "raw_suggestions.csv"),
              ["phrase", "raw_score", "hits"],
              [(s, round(score[s], 4), hits[s]) for s in sorted(score, key=lambda x: -score[x])])

    print(f"Keywords: {len(rows)}, rejected: {len(rejected)}, clusters: {len(clusters)}")
    if ok:
        print(f"Saved: {xlsx}")
    else:
        print("openpyxl is not installed: wrote CSV only. For XLSX: pip install openpyxl")
    print(f"Saved: {out_dir / f'{slug}_keywords.csv'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
